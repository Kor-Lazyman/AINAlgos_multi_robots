"""Preserve the authored report and append a Minor1 versus retry1 comparison."""
from pathlib import Path
import argparse, csv, hashlib, json

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/'report_final'
DEST=REPORT/'1.1.1e-retry1'
RECOVERY=REPORT/'recovery'
ORIGINAL=RECOVERY/'unsaved_word_draft_20261019.docx'
OUTPUT=RECOVERY/'combined_report_restored_with_minor1_comparison.docx'
NAMES=['Honeycomb','Spiral','Triangular truss','Motor inner','Twisted shell','Angle frame']
def fmt(seconds):
    n=round(float(seconds));d,n=divmod(n,86400);h,n=divmod(n,3600);m,s=divmod(n,60)
    return f'{d:02}:{h:02}:{m:02}:{s:02}'

def result_rows():
    minor=json.loads((REPORT/'minor1/stl_files/evaluation.json').read_text(encoding='utf-8'))
    minor_validation=json.loads((REPORT/'minor1/validation_data/evaluation.json').read_text(encoding='utf-8'))
    with (DEST/'results.csv').open(encoding='utf-8-sig',newline='') as stream:
        new=list(csv.DictReader(stream))
    together={}
    for split,prior in [('stl_files',minor),('validation_data',minor_validation)]:
        old=prior.get('jobs') or [dict(job,episodes=[episode],job_wall_time_s=episode['wall_time_s'])
                                   for job,episode in zip(prior['validation_jobs'],prior['episodes'])]
        current=[r for r in new if r['split']==split and
                 (split=='stl_files' or int(r['job'][:2]) < 6)]
        assert len(old)==len(current)==6
        rows=[]
        for i,(base,now) in enumerate(zip(old,current)):
            ep=base['episodes'][0]
            name=f'{i+1:02}' if split=='stl_files' else now['job']
            assert (base['id']==name or split=='stl_files') and ep['status']=='PASS'
            assert base['input_sha256']['target.stl']==json.loads(
                (ROOT/'output/train_e_retry1_40updates_20261019/training.json').read_text(encoding='utf-8'))[
                    'train_jobs' if split=='stl_files' else 'validation_jobs'][i]['target_sha256']
            old_time=float(ep['makespan_s']);new_time=float(now['makespan_s'])
            rows.append(dict(split=split,job=name,minor1_makespan_s=old_time,
                retry1_makespan_s=new_time,change_percent=(new_time/old_time-1)*100,
                minor1_computation_s=float(base['job_wall_time_s']),
                retry1_computation_s=float(now['computation_s']),
                minor1_validator=ep['status'],retry1_validator=now['validator_status'],
                original_stl_sha256=base['input_sha256']['target.stl'],
                minor1_checkpoint_sha256=minor['model_sha256'],
                retry1_checkpoint_sha256=now['checkpoint_sha256']))
        together[split]=rows
    return together

def add_table(doc, headings, rows):
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    table=doc.add_table(rows=1,cols=len(headings))
    table.style='Table Grid' if 'Table Grid' in doc.styles else None
    for cell,text in zip(table.rows[0].cells,headings):cell.text=text
    for row in rows:
        for cell,text in zip(table.add_row().cells,row):cell.text=str(text)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_before=Pt(2);para.paragraph_format.space_after=Pt(2)
                for run in para.runs:run.font.size=Pt(8.5)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:run.bold=True

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--capture',action='store_true')
    parser.add_argument('--report',action='store_true')
    args=parser.parse_args()
    if args.capture==args.report:parser.error('Choose exactly one of --capture or --report')
    if not ORIGINAL.is_file():raise RuntimeError('Authored report backup is missing')
    data=result_rows()
    base=REPORT/'minor1/stl_files'
    if args.capture:
        from package_retry_report import validator_gantt
        for i in range(6):
            job=base/f'job_{i:02}/episode_000'
            inputs=json.loads((job/'validation/validation_inputs.json').read_text(encoding='utf-8'))['inputs']
            for filename in ['config.yaml','target.stl','trajectory.csv']:
                assert hashlib.sha256((job/'input'/filename).read_bytes()).hexdigest()==inputs[filename]['sha256']
            validator_gantt(job)
            print('minor1 WAAM Gantt captured',i+1,flush=True)
        return
    from docx import Document
    from docx.shared import Inches
    from docx.shared import Pt
    from docx.enum.style import WD_STYLE_TYPE
    rows=data['stl_files']+data['validation_data']
    csv_path=REPORT/'minor1_retry1_comparison.csv'
    with csv_path.open('w',encoding='utf-8-sig',newline='') as stream:
        writer=csv.DictWriter(stream,fieldnames=list(rows[0]))
        writer.writeheader();writer.writerows(rows)
    doc=Document(ORIGINAL)
    for style_name in ['Heading 1','Heading 2']:
        if style_name not in doc.styles:
            style=doc.styles.add_style(style_name,WD_STYLE_TYPE.PARAGRAPH)
            style.font.name='맑은 고딕'
            style.font.size=Pt(16 if style_name=='Heading 1' else 12)
    doc.add_page_break()
    doc.add_heading('Minor1과 1.1.1.e retry1 결과 비교',1)
    doc.add_paragraph('동일한 원본 학습 STL 6개를 Minor1 최우수 체크포인트(1,024 step)와 '
                      '1.1.1.e retry1 최종 체크포인트(2,560 step)로 평가한 결과이다. '
                      '두 모델의 정책, STL 배치 및 학습량이 다르므로 표의 차이는 방법 전체의 결과 차이다.')
    doc.add_heading('학습 STL Makespan',2)
    add_table(doc,['STL','Minor1\nDD:HH:MM:SS','Retry1\nDD:HH:MM:SS','변화율','Validator\nMinor1 / Retry1'],[
        [r['job']+' '+NAMES[int(r['job'])-1],fmt(r['minor1_makespan_s']),
         fmt(r['retry1_makespan_s']),f'{r["change_percent"]:+.2f}%',
         r['minor1_validator']+' / '+r['retry1_validator']]
        for r in data['stl_files']])
    doc.add_paragraph('변화율 = (Retry1 Makespan / Minor1 Makespan − 1) × 100. 음수는 Retry1의 제작시간 감소를 뜻한다.')
    doc.add_heading('사례별 계산 시간',2)
    add_table(doc,['STL','Minor1\nDD:HH:MM:SS','Retry1\nDD:HH:MM:SS'],[
        [r['job'],fmt(r['minor1_computation_s']),fmt(r['retry1_computation_s'])]
        for r in data['stl_files']])
    doc.add_paragraph('기록된 계산시간의 측정 구간은 버전별 실행기에 따라 다르다. '
                      'Retry1의 값은 trajectory 생성 구간이며 WAAM Validator 실행시간을 제외한다. '
                      '정확한 초 단위 값은 minor1_retry1_comparison.csv에 있다.')
    doc.add_page_break()
    doc.add_heading('공통 validation 6개 Makespan',2)
    add_table(doc,['입력','Minor1\nDD:HH:MM:SS','Retry1\nDD:HH:MM:SS','변화율','Validator'],[
        [r['job'],fmt(r['minor1_makespan_s']),fmt(r['retry1_makespan_s']),
         f'{r["change_percent"]:+.2f}%',r['minor1_validator']+' / '+r['retry1_validator']]
        for r in data['validation_data']])
    doc.add_heading('작성자 설명',2);doc.add_paragraph(' ')
    for i,row in enumerate(data['stl_files']):
        doc.add_page_break()
        label=row['job']+' '+NAMES[i]
        doc.add_heading('Minor1과 Retry1 WAAM Validator 간트 비교',1)
        doc.add_paragraph(label+'  |  Minor1 '+fmt(row['minor1_makespan_s'])+
                          '  |  Retry1 '+fmt(row['retry1_makespan_s']))
        doc.add_paragraph('Minor1')
        doc.add_picture(str(base/f'job_{i:02}/episode_000/waam_validator_gantt.png'),width=Inches(6.45))
        doc.add_paragraph('1.1.1.e retry1')
        doc.add_picture(str(DEST/'stl_files'/row['job']/'waam_validator_gantt.png'),width=Inches(6.45))
    doc.save(OUTPUT)
    source=Document(ORIGINAL);result=Document(OUTPUT)
    assert len(result.inline_shapes)==len(source.inline_shapes)+12
    assert len(result.tables)==len(source.tables)+3
    for before,after in zip(source.paragraphs,result.paragraphs):assert before.text==after.text
    for before,after in zip(source.tables,result.tables):
        for b_row,a_row in zip(before.rows,after.rows):
            assert [cell.text for cell in b_row.cells]==[cell.text for cell in a_row.cells]
    print('Preserved authored content; added comparison to',OUTPUT,flush=True)

if __name__=='__main__':main()
