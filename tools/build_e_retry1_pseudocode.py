from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'C:\Users\lazzyman\Desktop\AINalgos2\version 1.1.1.e')
REPORT = ROOT / 'report_final'
OUT = REPORT / '1.1.1e-retry1' / 'pseudocode'
OUT.mkdir(parents=True, exist_ok=True)

lines = [
    ('Input:', 'STL M; robot configuration C; train set Dtr; validation set Dval'),
    ('Output:', 'trained PPO policy π and collision-checked trajectories'),
    ('Preprocessing and layer environment', ''),
    ('1', 'Sample interior solid points from M; assign weight 1 to every 3D sample.'),
    ('2', 'Project samples to XY and retain multiplicity from different Z heights.'),
    ('3', 'Search feasible rigid XY offsets; reject placements outside workspace or reach union.'),
    ('4', 'For each feasible offset, assign each projected point to its nearest reachable robot.'),
    ('5', 'Choose the offset minimizing the standard deviation of the three point shares;'),
    ('', 'break ties by maximum share error, travel cost, then offset magnitude.'),
    ('6', 'Apply the chosen single XY translation to the whole STL; slice the placed mesh.'),
    ('7', 'For each layer, generate fixed-spacing hatch segments at the selected hatch angle.'),
    ('8', 'Split segments at reach-circle boundaries and nearest-base bisectors;'),
    ('', 'assign every resulting segment to its nearest reachable robot.'),
    ('PPO action and path construction', ''),
    ('9', 'Policy selects hatch angle (8), each robot start XY (33 × 33 grid),'),
    ('', 'each sweep direction (CW/CCW), and one of 6 layer-priority permutations.'),
    ('10', 'Project each requested start onto that robot’s owned segment; split there,'),
    ('', 'then order its segments by the selected angular sweep.'),
    ('Collision retry and scheduling', ''),
    ('11', 'Visit robots in priority order. For each owned task, test approach, deposition,'),
    ('', 'and provisional safe return against other robots’ reserved trajectories.'),
    ('12', 'If safe, reserve the task. If blocked, keep it pending and test other owned tasks.'),
    ('13', 'Repeat a pending sweep after any task is accepted; stop when none remain blocked'),
    ('', 'or a sweep accepts no task.'),
    ('14', 'Transfer only unresolved tasks to the highest-priority reachable robot;'),
    ('', 'if none qualifies, retain them in the original robot’s deferred queue.'),
    ('15', 'Append deferred work after primary routes; delay its release from home only'),
    ('', 'when required for collision clearance. Preserve every task and deposition length.'),
    ('16', 'Check task coverage, physical collision clearance, and trajectory semantics.'),
    ('Training and evaluation', ''),
    ('17', 'Set layer reward rℓ = −Tℓ; episode return = −Σℓ Tℓ = −makespan.'),
    ('18', 'Update PPO for 40 updates × 64 steps; evaluate periodically and retain the'),
    ('', 'checkpoint with the lowest validation mean makespan. No early stopping.'),
    ('19', 'Run the official WAAM Validator for evaluation diagnostics only; it is not reward.'),
]

md = '''# 1.1.1.e retry1 수도코드\n\n**Algorithm 1. XY 투영 균등배치와 충돌 작업 재검토를 결합한 레이어 단위 PPO 스케줄링**\n\n```text\nInput: STL M, 로봇 설정 C, 학습 데이터 Dtr, 검증 데이터 Dval\nOutput: 학습된 PPO 정책과 충돌 검사를 마친 로봇 궤적\n\n# 전처리와 레이어 환경 구성\n각 3D 내부 샘플 점에 가중치 1을 부여하고 XY로 사영한다.\n서로 다른 Z에서 같은 XY로 사영된 점은 중복도(점 개수)를 유지한다.\n작업공간과 세 로봇 도달영역에 모두 들어가는 강체 XY 이동 후보를 탐색한다.\n각 후보에서 점을 가장 가까운 도달 가능 로봇에 배정한다.\n세 로봇의 점 비율 표준편차를 최소화하고, 동률이면 최대 비율 오차,\n이동 비용, 이동량 순으로 결정한다. 이 하나의 XY 이동을 STL 전체에 적용한다.\n배치된 STL을 슬라이싱하고 레이어별 고정 비드 간격 해치 선분을 생성한다.\n도달 원 경계와 최근접 로봇 경계에서 선분을 분할하고, 모든 선분을\n가장 가까운 도달 가능 로봇에 소유시킨다.\n\n# PPO action과 경로 구성\n각 레이어에서 PPO가 해치 각도(8), 로봇별 시작 XY(33×33), 회전 방향(CW/CCW),\n레이어 우선순위 순열(6개)을 선택한다.\n요청된 시작점을 소유 선분 위로 투영해 선분을 나누고 각도 순으로 정렬한다.\n\n# 충돌 재검토와 스케줄링\n우선순위 순으로 로봇을 방문한다. 각 소유 작업의 접근·적층·안전 복귀가\n다른 로봇의 예약 궤적과 충돌하지 않는지 검사한다.\n안전하면 예약한다. 막히면 해당 작업을 보류하고 나머지 소유 작업을 계속 검사한다.\n작업을 하나라도 예약한 뒤 보류 작업을 재검사한다. 보류가 없거나, 한 번의\n검사에서 새 작업을 예약하지 못하면 재검토를 끝낸다.\n남은 작업만 도달 가능한 상위 우선순위 로봇에 넘긴다. 대상이 없으면 원래\n로봇의 지연 큐에 보관한다. 지연 작업은 주 경로 뒤에 붙이며, 충돌 회피가\n필요할 때만 홈 위치에서 출발을 늦춘다. 작업 누락·중복, 적층 길이, 충돌,\n궤적 의미를 확인한다.\n\n# 학습과 평가\n레이어 보상 r_l = -T_l; 에피소드 보상 = -Σ_l T_l = -makespan.\nPPO를 40 updates × 64 steps 학습하고 검증 평균 makespan이 가장 낮은 체크포인트를 보존한다.\nEarly stopping은 사용하지 않는다. WAAM Validator는 평가 진단에만 사용하고 보상에는 넣지 않는다.\n```\n\n**학습 설정:** PPO MlpPolicy, pi/vf 은닉층 [128, 128], learning rate 0.0003, γ=1.0, GAE λ=0.95, batch size 32, epoch 10, 총 2,560 steps. 정책 계산 CUDA, geometry CPU. 초기 가중치는 무작위이며 기존 체크포인트를 불러오지 않았다.\n\n**구현 근거:** `runtime/positioning.py`, `runtime/layer_assignment_env.py`, `runtime/handoff_scheduler.py`, `runtime/train.py`, `model/training.json`.\n'''
(OUT / 'algorithm.md').write_text(md, encoding='utf-8')
(OUT / 'algorithm.txt').write_text('\n'.join((f'{n:>2}  {t}' if n else f'    {t}') for n, t in lines), encoding='utf-8')

# Paper-style single-page algorithm figure.
W, H = 2200, 3000
img = Image.new('RGB', (W, H), 'white')
draw = ImageDraw.Draw(img)
font_path = r'C:\Windows\Fonts\times.ttf'
bold_path = r'C:\Windows\Fonts\timesbd.ttf'
title = ImageFont.truetype(bold_path, 58)
subtitle = ImageFont.truetype(font_path, 30)
body = ImageFont.truetype(font_path, 33)
section = ImageFont.truetype(bold_path, 37)
numfont = ImageFont.truetype(bold_path, 30)
draw.text((145, 100), 'Algorithm 1', font=title, fill='#111111')
draw.text((145, 180), 'Projected placement and retry-before-handoff layer scheduling', font=subtitle, fill='#333333')
draw.line((145, 235, W-145, 235), fill='#222222', width=3)
y = 275
line_h = 61
for label, text in lines:
    if not label:
        draw.text((250, y), text, font=body, fill='#111111')
        y += line_h
    elif label in ('Preprocessing and layer environment', 'PPO action and path construction', 'Collision retry and scheduling', 'Training and evaluation'):
        y += 10
        draw.text((145, y), label, font=section, fill='#111111')
        y += 58
    elif label in ('Input:', 'Output:'):
        f = ImageFont.truetype(bold_path, 32)
        draw.text((175, y), label, font=f, fill='#111111')
        draw.text((330, y), text, font=body, fill='#111111')
        y += line_h
    else:
        draw.text((165, y), label, font=numfont, fill='#444444')
        draw.text((250, y), text, font=body, fill='#111111')
        y += line_h
draw.line((145, y+12, W-145, y+12), fill='#222222', width=3)
foot = 'Reward: rℓ = −Tℓ; episode return = −makespan     |     Validator: evaluation only'
draw.text((145, y+35), foot, font=subtitle, fill='#222222')
img.save(OUT / 'algorithm.png', dpi=(300, 300))

# Append to a candidate copy of the current restored report; source remains untouched.
source = REPORT / 'combined_report.docx'
candidate = REPORT / 'recovery' / 'combined_report_with_e_retry1_pseudocode_v2.docx'
candidate.parent.mkdir(parents=True, exist_ok=True)
doc = Document(source)
doc.add_page_break()
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.add_run().add_picture(str(OUT / 'algorithm.png'), width=Inches(6.45))
p2.paragraph_format.space_after = Pt(0)
doc.save(candidate)
print(candidate)
