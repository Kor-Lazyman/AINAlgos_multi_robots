"""Create a combined evidence report from the frozen result package."""
from pathlib import Path
from copy import deepcopy
import json,hashlib
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from package_report import DEST,OLD,NAMES,fmt

ALGORITHMS={
'minor2':[
'Input: STL mesh M, robot configuration C, PPO policy pi',
'Output: three collision-checked robot trajectories',
'Slice M into target polygons P[1], ..., P[L]',
'for each layer l = 1, ..., L do',
'    Observe layer geometry and current episode state',
'    a <- pi(s): angle, spacing, corner, reversal,',
'               partition axis, quota, robot permutation',
'    Generate hatch segments inside P[l]',
'    Order segments using the corner and reversal actions',
'    Partition spatial bands by deposition-length quotas',
'    Reassign unreachable segments to a reachable robot',
'    Construct concurrent routes including home returns',
'    if the concurrent routes contain a collision then',
'        Rebuild this entire layer as sequential routes',
'    Reject infeasible paths; otherwise append this layer',
'    Compute layer IoU, duration T and reference time B',
'    u <- mean robot deposition time / T',
'    if IoU > 0.9 then reward <- 10 * u * B / T',
'    else reward <- terminal Validator outcome, or 0',
'    Store transition (s, a, reward, next state)',
'end for',
'During training: update PPO from rollout transitions',
'Validate the full trajectory and export results',
],
'1.1.1e':[
'Input: STL M, robot configuration C, PPO policy pi',
'Output: three collision-checked robot trajectories',
'Sample equal-weight interior points from M',
'Project to XY, preserving the multiplicity along Z',
'Search one feasible XY shift balancing nearest owners',
'Translate the entire STL once; slice into target layers',
'for each target layer do',
'    a <- pi(s): hatch angle, each robot start XY and',
'               sweep direction, layer priority order',
'    Generate hatch paths; partition by nearest reach',
'    Project start XY onto each robot owned segments',
'    Split at the selected point; order each robot path',
'    for robot r in priority order do',
'        for each owned task q in its path order do',
'            Check approach, deposition and safe return',
'            if safe then reserve q for r',
'            else queue q for the highest reachable',
'                 higher-priority robot; keep r if none',
'        Continue checking the remaining owned tasks',
'    Append deferred queues after primary routes',
'    Delay queue release at home only when required',
'    Verify every task is deposited exactly once',
'    Check physical collisions and trajectory semantics',
'    reward <- negative layer makespan; append layer',
'end for; train PPO with gamma = 1',
'Export trajectories; run Validator for evaluation only',
]}

def pseudocode(v):
    out=DEST/v/'pseudocode';out.mkdir(exist_ok=True)
    lines=ALGORITHMS[v];w=1900;h=2450
    im=Image.new('RGB',(w,h),'white');d=ImageDraw.Draw(im)
    title=ImageFont.truetype('C:/Windows/Fonts/timesbd.ttf',56)
    font=ImageFont.truetype('C:/Windows/Fonts/times.ttf',42)
    number=ImageFont.truetype('C:/Windows/Fonts/times.ttf',32)
    d.line((65,65,w-65,65),fill='black',width=4)
    d.text((80,85),f'Algorithm {1 if v=="minor2" else 2}   {v} layer scheduling',font=title,fill='black')
    d.line((65,170,w-65,170),fill='black',width=3)
    for i,line in enumerate(lines):
        y=220+i*78;d.text((80,y+6),str(i+1),font=number,fill='#555')
        d.text((145,y),line,font=font,fill='black')
        assert d.textbbox((145,y),line,font=font)[2]<w-50,line
    d.line((65,2350,w-65,2350),fill='black',width=3)
    im.save(out/'algorithm.png');(out/'algorithm.txt').write_text('\n'.join(lines),encoding='utf-8')
    return out/'algorithm.png'

def main():
    data=json.loads((DEST/'report_data.json').read_text(encoding='utf-8'))
    doc=Document(OLD/'INE4107_combined_report.docx')
    metadata=deepcopy(doc.tables[0]._element)
    for el in list(doc._element.body):
        if el.tag!=qn('w:sectPr'):doc._element.body.remove(el)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        s=doc.styles[name];s.font.name='맑은 고딕';s.font.color.rgb=RGBColor(0,0,0)
        s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(21)
    doc.styles['Heading 1'].font.size=Pt(16)
    doc.styles['Heading 2'].font.size=Pt(12)
    def heading(t):doc.add_heading(t,1)
    def page(t):doc.add_page_break();heading(t)
    def table(headers,rows,size=9):
        t=doc.add_table(rows=1,cols=len(headers))
        if 'Table Grid' in doc.styles:t.style='Table Grid'
        else:
            borders=OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'999999');borders.append(el)
            t._tbl.tblPr.append(borders)
        for c,s in zip(t.rows[0].cells,headers):c.text=s
        for row in rows:
            for c,s in zip(t.add_row().cells,row):c.text=str(s)
        for row in t.rows:
            trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    for r in p.runs:r.font.size=Pt(size)
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for c in t.rows[0].cells:
            for r in c.paragraphs[0].runs:r.bold=True
        return t
    doc.add_paragraph('Minor2 및 1.1.1e 실험 결과',style='Title')
    doc._element.body.insert(len(doc._element.body)-1,metadata)
    doc.add_paragraph('학습 STL과 validation 데이터에 대한 최종 모델의 makespan, 계산시간, Validator 결과와 로봇 활용률을 정리하였다.')
    doc.add_paragraph('두 실험 모두 2,560-step 정책을 사용한다. 1.1.1.e는 40-update 최종 모델이며, 충돌 작업 재검토 방식으로 수정하기 전 결과이다.')
    table(['구분','Minor2','1.1.1.e'],[
        ['학습 STL 수','6','6'],['Validation 수','6','8'],['평가 정책','최종 2,560 step','최종 2,560 step'],
        ['학습 STL makespan 합계',fmt(sum(r['makespan_s'] for r in data if r['version']=='minor2' and r['split']=='stl_files')),fmt(sum(r['makespan_s'] for r in data if r['version']=='1.1.1e' and r['split']=='stl_files'))],
        ['공통 validation 6개 합계',fmt(sum(r['makespan_s'] for r in data if r['version']=='minor2' and r['split']=='validation_data')),fmt(sum(r['makespan_s'] for r in data if r['version']=='1.1.1e' and r['split']=='validation_data' and int(r['job'][:2])<6))],
        ['Validator','12 / 12 PASS','14 / 14 PASS']])
    doc.add_paragraph('시간 표기 DD:HH:MM:SS는 초 단위 반올림이다. 정확한 값은 all_results.csv에 수록하였다. 활용률은 (이동시간 + 적층시간) / makespan이며 동시 적층률과 다르다.')
    doc.add_paragraph('계산시간은 저장된 실행 기록의 측정값으로 캐시와 측정 범위가 다를 수 있다. e의 추가 validation 2개는 공통 비교 합계에서 제외하였다.')
    heading('코멘트');doc.add_paragraph(' ')
    for v in ['minor2','1.1.1e']:
        for split,label in [('stl_files','학습 STL'),('validation_data','Validation')]:
            records=[r for r in data if r['version']==v and r['split']==split]
            page(f'{v} {label} 적용 결과')
            table(['모델명','Makespan\nDD:HH:MM:SS','계산시간\nDD:HH:MM:SS','Validator'],[
                [(r['job']+' '+NAMES[int(r['job'])-1]) if split=='stl_files' else r['job'],fmt(r['makespan_s']),fmt(r['computation_s']),r['validator_status']] for r in records],8.5)
            doc.add_heading('로봇별 활용률',2)
            table(['모델','Robot 1','Robot 2','Robot 3','레이어 수'],[[r['job'],*[f'{r[f"robot{i}_utilization"]*100:.2f}%' for i in (1,2,3)],r['layers']] for r in records])
            doc.add_heading('코멘트',2);doc.add_paragraph(' ')
            for n in range(0,len(records),3):
                page(f'{v} {label} 간트차트')
                for r in records[n:n+3]:
                    p=doc.add_paragraph(f'{r["job"]}   Makespan {fmt(r["makespan_s"])}   Validator {r["validator_status"]}')
                    p.paragraph_format.keep_with_next=True
                    doc.add_picture(str(DEST/r['artifact_dir']/'gantt.png'),width=Inches(6.45))
        page(f'{v} 학습 설정')
        tr=json.loads((DEST/v/'model/training.json').read_text(encoding='utf-8'))
        if v=='minor2':
            p=tr['ppo'];rows=[('PPO rollout',p['n_steps']),('batch_size',p['batch_size']),('learning_rate',p['learning_rate']),('gamma',p['gamma']),('gae_lambda',p['gae_lambda']),('seed',p['seed']),('device',p['device']),('geometry','CPU'),('학습 업데이트',40),('학습 step',2560),('action','MultiDiscrete [8, 3, 4, 2, 2, 3, 6]'),('위치 조정 네트워크','없음'),('보상','IoU > 0.9: 10 × 평균 적층 활용률 × 기준시간 / 레이어 시간'),('기타 보상','IoU ≤ 0.9: terminal PASS +1, FAIL/ERROR −1, 그 외 0; invalid −2')]
        else:
            a=tr['arguments'];rows=[('PPO rollout',a['n_steps']),('batch_size',a['batch_size']),('epochs',a['epochs']),('learning_rate',tr['learning_rate']),('gamma',tr['gamma']),('gae_lambda',tr['gae_lambda']),('seed',a['seed']),('device',a['device']),('geometry','CPU'),('학습 업데이트',40),('학습 step',2560),('action','MultiDiscrete [8,33,33,2,33,33,2,33,33,2,6]'),('위치 조정','동일 가중치 사영 점의 비중을 맞추는 XY 이동'),('보상','−레이어 makespan'),('early stopping','없음'),('최종 / 최우수 정책','2,560 / 2,048 step; 보고서는 최종 정책 사용')]
        table(['설정','값'],rows,9)
        doc.add_heading('이 알고리즘을 적용한 이유',2);doc.add_paragraph(' ')
        doc.add_heading('특징과 장단점',2);doc.add_paragraph(' ')
    for v in ['minor2','1.1.1e']:
        p=doc.add_paragraph();p.paragraph_format.page_break_before=True
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
        p.add_run().add_picture(str(pseudocode(v)),width=Inches(6.35))
    out=DEST/'combined_report.docx';doc.save(out)
    print(out)
if __name__=='__main__':main()
