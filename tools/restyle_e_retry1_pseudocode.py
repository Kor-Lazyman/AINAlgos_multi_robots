from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(r'C:\Users\lazzyman\Desktop\AINalgos2\version 1.1.1.e')
REPORT = ROOT / 'report_final'
PSEUDO = REPORT / '1.1.1e-retry1' / 'pseudocode'
IMAGE = PSEUDO / 'algorithm_paper_style.png'

rows = [
    ('Input:', 'STL M, robot configuration C, training set D_tr, validation set D_val'),
    ('Output:', 'trained PPO policy and collision-checked robot trajectories'),
    ('01', 'Sample solid interior points from M; give every 3D sample weight 1'),
    ('02', 'Project samples to XY and preserve multiplicity from different Z heights'),
    ('03', 'Search rigid XY shifts that fit the workspace and all robot reach regions'),
    ('04', 'Assign each projected point to its nearest reachable robot'),
    ('05', 'Choose the shift minimizing the standard deviation of the three point shares;'),
    ('', 'break ties by maximum share error, travel cost, then shift magnitude'),
    ('06', 'Translate the entire STL by this one XY shift, then slice it into layers'),
    ('07', 'For each layer, create fixed-spacing hatch segments at the selected angle'),
    ('08', 'Split segments at reach-circle boundaries and nearest-base bisectors;'),
    ('', 'assign each segment to its nearest reachable robot'),
    ('09', 'Policy selects hatch angle (8), each robot start XY (33 x 33),'),
    ('', 'each sweep direction (CW/CCW), and one of 6 priority orders'),
    ('10', 'Project each requested start onto an owned segment, split it there, and order'),
    ('', 'the robot segments by the selected angular sweep'),
    ('11', 'for each robot r in the selected priority order do'),
    ('12', '    for each deposition task q owned by r do'),
    ('13', '        Test approach, deposition, and safe return against reserved trajectories'),
    ('14', '        if safe then reserve q for r'),
    ('15', '        else keep q pending and continue with the remaining owned tasks'),
    ('16', '    Retry pending tasks after any task is accepted; stop if none remain or'),
    ('', '    a full sweep accepts no task'),
    ('17', 'Transfer only unresolved tasks to the highest-priority robot that can reach them;'),
    ('', 'if none qualifies, retain the tasks in the original robot deferred queue'),
    ('18', 'Append deferred work after primary routes; delay its home release only if needed'),
    ('19', 'Verify every task occurs once, deposition length is preserved, and trajectories'),
    ('', 'pass physical collision and trajectory-semantics checks'),
    ('20', 'Set layer reward r_l = -T_l; episode return = -sum_l T_l = -makespan'),
    ('21', 'Train PPO for 40 updates x 64 steps; keep the lowest validation-mean-makespan'),
    ('', 'checkpoint; do not early-stop. Run WAAM Validator for evaluation only'),
]

markdown = '''# Algorithm 1. 1.1.1.e retry1 layer PPO scheduling\n\n**Input:** STL M, robot configuration C, training set D_tr, validation set D_val\n\n**Output:** trained PPO policy and collision-checked robot trajectories\n\n```text\n01  Sample solid interior points from M; give every 3D sample weight 1\n02  Project samples to XY and preserve multiplicity from different Z heights\n03  Search rigid XY shifts that fit the workspace and all robot reach regions\n04  Assign each projected point to its nearest reachable robot\n05  Choose the shift minimizing the standard deviation of the three point shares;\n    break ties by maximum share error, travel cost, then shift magnitude\n06  Translate the entire STL by this one XY shift, then slice it into layers\n07  For each layer, create fixed-spacing hatch segments at the selected angle\n08  Split segments at reach-circle boundaries and nearest-base bisectors;\n    assign each segment to its nearest reachable robot\n09  Policy selects hatch angle (8), each robot start XY (33 x 33),\n    each sweep direction (CW/CCW), and one of 6 priority orders\n10  Project each requested start onto an owned segment, split it there, and order\n    the robot segments by the selected angular sweep\n11  for each robot r in the selected priority order do\n12      for each deposition task q owned by r do\n13          Test approach, deposition, and safe return against reserved trajectories\n14          if safe then reserve q for r\n15          else keep q pending and continue with the remaining owned tasks\n16      Retry pending tasks after any task is accepted; stop if none remain or\n        a full sweep accepts no task\n17  Transfer only unresolved tasks to the highest-priority robot that can reach them;\n    if none qualifies, retain the tasks in the original robot deferred queue\n18  Append deferred work after primary routes; delay its home release only if needed\n19  Verify every task occurs once, deposition length is preserved, and trajectories\n    pass physical collision and trajectory-semantics checks\n20  Set layer reward r_l = -T_l; episode return = -sum_l T_l = -makespan\n21  Train PPO for 40 updates x 64 steps; keep the lowest validation-mean-makespan\n    checkpoint; do not early-stop. Run WAAM Validator for evaluation only\n```\n\n**Training:** PPO MlpPolicy; policy/value MLP layers [128, 128]; learning rate 0.0003; gamma 1.0; GAE lambda 0.95; batch size 32; 10 epochs; 2,560 steps. Policy runs on CUDA and geometry on CPU. Training starts from fresh random weights.\n'''
(PSEUDO / 'algorithm.md').write_text(markdown, encoding='utf-8')
(PSEUDO / 'algorithm.txt').write_text(markdown.split('```text\n', 1)[1].split('\n```', 1)[0] + '\n', encoding='utf-8')

# Reference-like paper algorithm image: pale blue title band, blue title and line numbers,
# monospaced body, restrained spacing, and numbered pseudocode.
W, H = 2200, 2600
im = Image.new('RGB', (W, H), '#FFFFFF')
d = ImageDraw.Draw(im)
d.rectangle((0, 0, W, 185), fill='#E8F0FA')
title_font = ImageFont.truetype(r'C:\Windows\Fonts\consolab.ttf', 55)
body_font = ImageFont.truetype(r'C:\Windows\Fonts\consola.ttf', 36)
num_font = ImageFont.truetype(r'C:\Windows\Fonts\consola.ttf', 34)
d.text((62, 48), 'Algorithm 1. 1.1.1.e retry1 - layer PPO scheduling',
       font=title_font, fill='#245D91')
y, line_h = 230, 68
number_x, text_x, max_x = 62, 220, W - 62
for label, content in rows:
    font = body_font
    if label:
        d.text((number_x, y), label, font=num_font, fill='#91A5B9')
    d.text((text_x, y), content, font=font, fill='#253449')
    y += line_h
if y + 30 > H:
    raise RuntimeError(f'Image overflow: final y={y}, image height={H}')
im.save(IMAGE, dpi=(300, 300))

# Replace only the final pseudocode image in the current report.
source = REPORT / 'combined_report.docx'
candidate = REPORT / 'recovery' / 'combined_report_with_e_retry1_pseudocode_styled_v2.docx'
doc = Document(source)
image_paragraph = next((p for p in reversed(doc.paragraphs)
                        if p._p.xpath('.//w:drawing')), None)
if image_paragraph is None:
    raise RuntimeError('Could not find the existing final pseudocode image')
parent = image_paragraph._element.getparent()
parent.remove(image_paragraph._element)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
p.add_run().add_picture(str(IMAGE), width=Inches(6.45))
doc.save(candidate)
print(candidate)
