"""Append the final retry1 results and Validator Gantt captures to the report."""
from pathlib import Path
import csv,json
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/'report_final'
VARIANT=REPORT/'1.1.1e-retry1'
SOURCE=REPORT/'combined_report.docx'
OUTPUT=REPORT/'combined_report_with_retry1.docx'
NAMES=['Honeycomb','Spiral','Triangular truss','Motor inner','Twisted shell','Angle frame']
def fmt(s):
    n=round(float(s));d,n=divmod(n,86400);h,n=divmod(n,3600);m,s=divmod(n,60)
    return f'{d:02}:{h:02}:{m:02}:{s:02}'
def add_table(doc,heads,rows):
    tab=doc.add_table(rows=1,cols=len(heads))
    if 'Table Grid' in doc.styles:tab.style='Table Grid'
    for c,value in zip(tab.rows[0].cells,heads):c.text=value
    for row in rows:
        for c,value in zip(tab.add_row().cells,row):c.text=str(value)
    for row in tab.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.space_after=Pt(2)
                for run in p.runs:run.font.size=Pt(9)
    tab.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for c in tab.rows[0].cells:
        for run in c.paragraphs[0].runs:run.bold=True
def main():
    train=json.loads((ROOT/'output/train_e_retry1_40updates_20261019/training.json').read_text(encoding='utf-8'))
    if train['status']!='complete' or train['completed_updates']!=40:raise RuntimeError('Training incomplete')
    with (VARIANT/'results.csv').open(encoding='utf-8-sig',newline='') as stream:
        items=list(csv.DictReader(stream))
    stl=[r for r in items if r['split']=='stl_files']
    val=[r for r in items if r['split']=='validation_data']
    if len(stl)!=6 or len(val)!=8 or any(r['validator_status']!='PASS' for r in items):
        raise RuntimeError('Incomplete result package')
    doc=Document(SOURCE)
    doc.add_page_break()
    doc.add_heading('1.1.1.e retry1 최종 학습 결과',1)
    doc.add_paragraph('40 updates, 2,560 step 학습의 최종 체크포인트를 학습 STL 6개와 validation 8개에 적용한 결과이다. 학습 시간은 '
                      +fmt(train['wall_s'])+'이며 전처리, 학습, validation 및 최종 평가를 포함한다.')
    doc.add_heading('학습 STL 모델별 적용 결과',2)
    add_table(doc,['모델명','Makespan\n(DD:HH:MM:SS)','계산 시간\n(DD:HH:MM:SS)','Validator 적용 결과'],[
        [r['job']+' ('+NAMES[int(r['job'])-1]+')',fmt(r['makespan_s']),fmt(r['computation_s']),r['validator_status']]
        for r in stl])
    doc.add_paragraph('계산 시간은 최종 모델로 해당 STL의 trajectory를 생성한 경과시간이며, 별도 WAAM Validator 실행 시간은 포함하지 않는다. '
                      '학습 전체 시간과 다른 수치이며, 정확한 초 단위 값은 results.csv에 기록하였다.')
    doc.add_heading('Validator 결과',2)
    doc.add_paragraph('학습 STL 6개와 validation 8개 모두 PASS, 충돌 이벤트 0건. '
                      '아래 간트 이미지는 각 PASS 판정에 사용된 trajectory.csv를 WAAM Validator 대시보드의 '
                      '간트 함수로 표시한 뒤 브라우저에서 캡처했다.')
    doc.add_heading('코멘트',2);doc.add_paragraph(' ')
    for chunk in (stl[:3],stl[3:]):
        doc.add_page_break();doc.add_heading('1.1.1.e retry1 WAAM Validator 간트차트',1)
        for row in chunk:
            name=row['job']+' ('+NAMES[int(row['job'])-1]+')'
            paragraph=doc.add_paragraph(name+'  |  '+fmt(row['makespan_s'])+'  |  Validator PASS')
            paragraph.paragraph_format.keep_with_next=True
            doc.add_picture(str(VARIANT/'stl_files'/row['job']/'waam_validator_gantt.png'),width=Inches(6.45))
    doc.add_page_break();doc.add_heading('1.1.1.e retry1 validation 결과',1)
    add_table(doc,['모델명','Makespan\n(DD:HH:MM:SS)','계산 시간\n(DD:HH:MM:SS)','Validator 적용 결과'],[
        [r['job'],fmt(r['makespan_s']),fmt(r['computation_s']),r['validator_status']] for r in val])
    doc.add_paragraph('각 validation의 Validator 간트 이미지와 HTML은 report_final/1.1.1e-retry1/validation_data에 보관했다.')
    doc.add_heading('코멘트',2);doc.add_paragraph(' ')
    doc.save(OUTPUT)
    print(OUTPUT)
if __name__=='__main__':main()
